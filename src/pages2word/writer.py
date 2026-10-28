"""Write the semantic model into a .docx that uses Word's own styles and lists."""

from __future__ import annotations

import copy
import io
import logging
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import XmlPart
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt, RGBColor
from docx.table import _Cell
from docx.text.paragraph import Paragraph as DocxParagraph
from lxml import etree

from .model import (
    Block,
    Document,
    Footnote,
    Image,
    PageBreak,
    Paragraph,
    Pencil,
    Run,
    Table,
    TextBox,
)
from .wordstyles import Profile, load_profile, save_profile

log = logging.getLogger(__name__)

W_NS = nsdecls("w")

# Word 365 defaults: Aptos 12 pt, 8 pt after, 1.15 lines; headings in accent blue.
STYLES_XML = f"""
<w:styles {W_NS}>
  <w:docDefaults>
    <w:rPrDefault><w:rPr>
      <w:rFonts w:asciiTheme="minorHAnsi" w:eastAsiaTheme="minorEastAsia" w:hAnsiTheme="minorHAnsi" w:cstheme="minorBidi"/>
      <w:kern w:val="2"/><w:sz w:val="24"/><w:szCs w:val="24"/>
      <w:lang w:val="{{lang}}" w:eastAsia="en-US" w:bidi="ar-SA"/>
    </w:rPr></w:rPrDefault>
    <w:pPrDefault><w:pPr><w:spacing w:after="160" w:line="278" w:lineRule="auto"/></w:pPr></w:pPrDefault>
  </w:docDefaults>
  <w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/><w:qFormat/></w:style>
  <w:style w:type="character" w:default="1" w:styleId="DefaultParagraphFont"><w:name w:val="Default Paragraph Font"/><w:uiPriority w:val="1"/><w:semiHidden/><w:unhideWhenUsed/></w:style>
  <w:style w:type="table" w:default="1" w:styleId="TableNormal"><w:name w:val="Normal Table"/><w:uiPriority w:val="99"/><w:semiHidden/><w:unhideWhenUsed/><w:tblPr><w:tblInd w:w="0" w:type="dxa"/><w:tblCellMar><w:top w:w="0" w:type="dxa"/><w:left w:w="108" w:type="dxa"/><w:bottom w:w="0" w:type="dxa"/><w:right w:w="108" w:type="dxa"/></w:tblCellMar></w:tblPr></w:style>
  <w:style w:type="numbering" w:default="1" w:styleId="NoList"><w:name w:val="No List"/><w:uiPriority w:val="99"/><w:semiHidden/><w:unhideWhenUsed/></w:style>
  {{headings}}
  <w:style w:type="paragraph" w:styleId="Title"><w:name w:val="Title"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:uiPriority w:val="10"/><w:qFormat/>
    <w:pPr><w:spacing w:after="80" w:line="240" w:lineRule="auto"/><w:contextualSpacing/></w:pPr>
    <w:rPr><w:rFonts w:asciiTheme="majorHAnsi" w:eastAsiaTheme="majorEastAsia" w:hAnsiTheme="majorHAnsi" w:cstheme="majorBidi"/><w:spacing w:val="-10"/><w:kern w:val="28"/><w:sz w:val="56"/><w:szCs w:val="56"/></w:rPr></w:style>
  {{extra}}
  <w:style w:type="paragraph" w:styleId="ListParagraph"><w:name w:val="List Paragraph"/><w:basedOn w:val="Normal"/><w:uiPriority w:val="34"/><w:qFormat/>
    <w:pPr><w:ind w:left="720"/><w:contextualSpacing/></w:pPr></w:style>
</w:styles>
"""

# Styles the converter needs beyond headings; added to user templates when missing.
EXTRA_STYLES = {
    "TableGrid": f"""<w:style {W_NS} w:type="table" w:styleId="TableGrid"><w:name w:val="Table Grid"/><w:basedOn w:val="TableNormal"/><w:uiPriority w:val="39"/>
    <w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>
    <w:tblPr><w:tblBorders><w:top w:val="single" w:sz="4" w:space="0" w:color="auto"/><w:left w:val="single" w:sz="4" w:space="0" w:color="auto"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="auto"/><w:right w:val="single" w:sz="4" w:space="0" w:color="auto"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="auto"/><w:insideV w:val="single" w:sz="4" w:space="0" w:color="auto"/></w:tblBorders></w:tblPr></w:style>""",
    "FootnoteText": f"""<w:style {W_NS} w:type="paragraph" w:styleId="FootnoteText"><w:name w:val="footnote text"/><w:basedOn w:val="Normal"/><w:uiPriority w:val="99"/><w:unhideWhenUsed/>
    <w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr><w:rPr><w:sz w:val="20"/><w:szCs w:val="20"/></w:rPr></w:style>""",
    "FootnoteReference": f"""<w:style {W_NS} w:type="character" w:styleId="FootnoteReference"><w:name w:val="footnote reference"/><w:uiPriority w:val="99"/><w:unhideWhenUsed/>
    <w:rPr><w:vertAlign w:val="superscript"/></w:rPr></w:style>""",
}

# (size half-points, before, after, major font?, italic?)
HEADING_SPECS = {
    1: (40, 360, 80, True, False),
    2: (32, 160, 80, True, False),
    3: (28, 160, 80, False, False),
    4: (24, 80, 40, False, True),
    5: (24, 80, 40, False, False),
    6: (24, 40, 0, False, True),
}


def _heading_style(level: int) -> str:
    size, before, after, major, italic = HEADING_SPECS[min(level, 6)]
    theme = "major" if major else "minor"
    return f"""
  <w:style w:type="paragraph" w:styleId="Heading{level}"><w:name w:val="heading {level}"/><w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:uiPriority w:val="9"/><w:qFormat/>
    <w:pPr><w:keepNext/><w:keepLines/><w:spacing w:before="{before}" w:after="{after}"/><w:outlineLvl w:val="{level - 1}"/></w:pPr>
    <w:rPr><w:rFonts w:asciiTheme="{theme}HAnsi" w:eastAsiaTheme="{theme}EastAsia" w:hAnsiTheme="{theme}HAnsi" w:cstheme="{theme}Bidi"/>
      {"<w:i/><w:iCs/>" if italic else ""}<w:color w:val="0F4761" w:themeColor="accent1" w:themeShade="BF"/><w:sz w:val="{size}"/><w:szCs w:val="{size}"/></w:rPr></w:style>"""


# Word's default bullet glyphs per level: filled dot, "o", filled square.
BULLETS = [("", "Symbol"), ("o", "Courier New"), ("", "Wingdings")]


def _bullet_abstract(abstract_id: int) -> str:
    levels = []
    for lvl in range(9):
        glyph, font = BULLETS[lvl % 3]
        left = 720 * (lvl + 1)
        levels.append(
            f'<w:lvl w:ilvl="{lvl}"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="{glyph}"/>'
            f'<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="{left}" w:hanging="360"/></w:pPr>'
            f'<w:rPr><w:rFonts w:ascii="{font}" w:hAnsi="{font}" w:hint="default"/></w:rPr></w:lvl>'
        )
    return (
        f'<w:abstractNum {W_NS} w:abstractNumId="{abstract_id}"><w:multiLevelType w:val="hybridMultilevel"/>'
        + "".join(levels)
        + "</w:abstractNum>"
    )


def _number_abstract(abstract_id: int, formats: dict[int, tuple[str, str]]) -> str:
    """formats: level -> (numFmt, suffix) taken from Pages; other levels use Word's 1. a. i. cycle."""
    cycle = ["decimal", "lowerLetter", "lowerRoman"]
    levels = []
    for lvl in range(9):
        fmt, suffix = formats.get(lvl, (cycle[lvl % 3], "."))
        text = f"({'%'}{lvl + 1})" if suffix == "()" else f"%{lvl + 1}{suffix}"
        roman = fmt == "lowerRoman"
        left = 720 * (lvl + 1)
        levels.append(
            f'<w:lvl w:ilvl="{lvl}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/><w:lvlText w:val="{text}"/>'
            f'<w:lvlJc w:val="{"right" if roman else "left"}"/>'
            f'<w:pPr><w:ind w:left="{left}" w:hanging="{180 if roman else 360}"/></w:pPr></w:lvl>'
        )
    return (
        f'<w:abstractNum {W_NS} w:abstractNumId="{abstract_id}"><w:multiLevelType w:val="hybridMultilevel"/>'
        + "".join(levels)
        + "</w:abstractNum>"
    )


def _copy_abstract(abstract_id: int, element) -> str:
    el = copy.deepcopy(element)
    el.set(qn("w:abstractNumId"), str(abstract_id))
    for tag in ("w:nsid", "w:tmpl", "w:numStyleLink", "w:styleLink"):
        for child in el.findall(qn(tag)):
            el.remove(child)
    return etree.tostring(el, encoding="unicode")


class _Numbering:
    def __init__(self, document, bullets=None):
        self.root = document.part.numbering_part.element
        for child in list(self.root):
            self.root.remove(child)
        self._abstract_count = 0
        self._num_count = 0
        if bullets is not None:  # the template's own bullet list definition
            self.bullet_abstract = self._abstract(_copy_abstract, bullets)
        else:
            self.bullet_abstract = self._abstract(_bullet_abstract)
        self.bullet_num = self._num(self.bullet_abstract)

    def _abstract(self, factory, *args) -> int:
        aid = self._abstract_count
        self._abstract_count += 1
        el = parse_xml(factory(aid, *args))
        # abstractNum elements must precede num elements
        nums = self.root.findall(qn("w:num"))
        if nums:
            nums[0].addprevious(el)
        else:
            self.root.append(el)
        return aid

    def _num(self, abstract_id: int, restart: bool = False) -> int:
        self._num_count += 1
        nid = self._num_count
        override = (
            '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride>'
            if restart
            else ""
        )
        self.root.append(
            parse_xml(
                f'<w:num {W_NS} w:numId="{nid}"><w:abstractNumId w:val="{abstract_id}"/>{override}</w:num>'
            )
        )
        return nid

    def new_numbered_list(self, formats: dict[int, tuple[str, str]]) -> int:
        return self._num(self._abstract(_number_abstract, formats), restart=True)


def _patch_theme_fonts(document):
    """Swap the template's Calibri/Cambria theme fonts for Word 365's Aptos."""
    for rel in document.part.rels.values():
        if rel.reltype.endswith("/theme"):
            part = rel.target_part
            xml = part.blob.decode("utf-8")
            xml = xml.replace('typeface="Cambria"', 'typeface="Aptos Display"').replace(
                'typeface="Calibri"', 'typeface="Aptos"'
            )
            part._blob = xml.encode("utf-8")


def _new_document(language: str, template: Path | None) -> tuple[object, Profile]:
    if template:
        document = docx.Document(str(template))
        profile = load_profile(document)
        body = document.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)
        root = document.styles.element
        present = {el.get(qn("w:styleId")) for el in root.findall(qn("w:style"))}
        for style_id, xml in EXTRA_STYLES.items():
            if style_id not in present:
                root.append(parse_xml(xml))
        return document, profile

    document = docx.Document()
    headings = "".join(_heading_style(n) for n in range(1, 7))
    styles = parse_xml(
        STYLES_XML.replace("{lang}", language)
        .replace("{headings}", headings)
        .replace(
            "{extra}", "".join(s.replace(f" {W_NS}", "") for s in EXTRA_STYLES.values())
        )
    )
    root = document.styles.element
    latent = root.find(qn("w:latentStyles"))
    for child in list(root):
        root.remove(child)
    root.append(styles.find(qn("w:docDefaults")))
    if latent is not None:
        root.append(latent)
    for style in styles.findall(qn("w:style")):
        root.append(style)
    _patch_theme_fonts(document)

    # A4 with Word's default 2.5 cm margins
    section = document.sections[0]
    section.page_width, section.page_height = docx.shared.Mm(210), docx.shared.Mm(297)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, docx.shared.Cm(2.5))
    return document, Profile()


def _set_lang(run, lang: str):
    rpr = run._r.get_or_add_rPr()
    el = rpr.find(qn("w:lang"))
    if el is None:
        el = parse_xml(f"<w:lang {W_NS}/>")
        rpr.append(el)
    el.set(qn("w:val"), lang)


def _set_shading(run, color: str):
    rpr = run._r.get_or_add_rPr()
    rpr.append(
        parse_xml(f'<w:shd {W_NS} w:val="clear" w:color="auto" w:fill="{color}"/>')
    )


def _set_numbering(paragraph, num_id: int, level: int):
    ppr = paragraph._p.get_or_add_pPr()
    ppr.append(
        parse_xml(
            f'<w:numPr {W_NS}><w:ilvl w:val="{level}"/><w:numId w:val="{num_id}"/></w:numPr>'
        )
    )


ALIGN = {
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

FOOTNOTES_CT = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
)
FOOTNOTE_SEPARATORS = f"""<w:footnotes {nsdecls("w", "r")}>
  <w:footnote w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr><w:r><w:separator/></w:r></w:p></w:footnote>
  <w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>
</w:footnotes>"""
EMU_PER_PT = 12700


class _Footnotes:
    """Creates word/footnotes.xml on first use (python-docx has no footnote API)."""

    def __init__(self, document):
        self.document = document
        self.root = None
        self.next_id = 1

    def add(self):
        if self.root is None:
            # a template's own footnotes part would clash with ours; its notes are gone anyway
            for r_id, rel in list(self.document.part.rels.items()):
                if rel.reltype == RT.FOOTNOTES:
                    self.document.part.drop_rel(r_id)
            self.root = parse_xml(FOOTNOTE_SEPARATORS)
            part = XmlPart(
                PackURI("/word/footnotes.xml"),
                FOOTNOTES_CT,
                self.root,
                self.document.part.package,
            )
            self.document.part.relate_to(part, RT.FOOTNOTES)
        note_id = self.next_id
        self.next_id += 1
        note = parse_xml(f'<w:footnote {W_NS} w:id="{note_id}"/>')
        self.root.append(note)
        return note_id, note


def _anchor_picture(run, dx: float, dy: float, z: int):
    """Turn the inline picture in `run` into a floating one next to the run's position."""
    inline = run._r.find(".//" + qn("wp:inline"))
    extent = inline.find(qn("wp:extent"))
    doc_pr = inline.find(qn("wp:docPr"))
    frame = inline.find(qn("wp:cNvGraphicFramePr"))
    graphic = inline.find(qn("a:graphic"))
    anchor = parse_xml(
        f'<wp:anchor {nsdecls("wp")} distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        f'relativeHeight="{251659264 + z}" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="character"><wp:posOffset>{round(dx * EMU_PER_PT)}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="line"><wp:posOffset>{round(dy * EMU_PER_PT)}</wp:posOffset></wp:positionV>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/></wp:anchor>'
    )
    anchor.insert(3, extent)
    for el in (doc_pr, frame, graphic):
        if el is not None:
            anchor.append(el)
    inline.getparent().replace(inline, anchor)


class WordWriter:
    def __init__(
        self,
        doc: Document,
        template: Path | None = None,
        empty_lines_to_page_break: int = 4,
    ):
        self.doc = doc
        self.language = doc.language or "cs-CZ"
        self.document, self.profile = _new_document(self.language, template)
        self.numbering = _Numbering(self.document, self.profile.bullets)
        if template:  # outputs carry the roles too, so they work as style documents
            save_profile(self.document, self.profile, self.numbering.bullet_abstract)
        self.footnotes = _Footnotes(self.document)
        self.empty_limit = empty_lines_to_page_break
        # With a user template its heading styles decide colours, not Pages.
        self.heading_colors = template is None
        section = self.document.sections[0]
        self.max_image_width = (
            section.page_width - section.left_margin - section.right_margin
        ) / EMU_PER_PT
        self._number_list: int | None = None
        self._number_formats: dict[int, tuple[str, str]] = {}
        self._fresh_cells: set = set()  # cells whose initial empty paragraph is unused
        self._z = 0

    def write(self, out: Path):
        self._blocks(self.document, self._collapse_empty_runs(self.doc.blocks))
        self.document.save(str(out))

    def _collapse_empty_runs(self, blocks: list[Block]) -> list[Block]:
        """Long runs of blank lines were used to push text to a new page; use a page break instead."""
        if not self.empty_limit:
            return blocks
        out: list[Block] = []
        blanks: list[Block] = []
        for b in blocks:
            if isinstance(b, Paragraph) and b.is_empty and not b.list_kind:
                blanks.append(b)
                continue
            if len(blanks) >= self.empty_limit and out and isinstance(b, Paragraph):
                b.page_break_before = True
            else:
                out.extend(blanks)
            blanks = []
            out.append(b)
        if len(blanks) < self.empty_limit:
            out.extend(blanks)
        return out

    # ---- blocks -------------------------------------------------------------------

    def _blocks(self, container, blocks: list[Block]):
        previous = None
        for b in blocks:
            if isinstance(b, Table | TextBox) and isinstance(previous, Table | TextBox):
                self._new_paragraph(container)  # Word would merge touching tables
            if isinstance(b, Paragraph):
                self._paragraph(container, b)
            elif isinstance(b, Table):
                self._table(container, b)
            elif isinstance(b, TextBox):
                self._text_box(container, b)
            previous = b

    def _new_paragraph(self, container):
        if isinstance(container, _Cell) and container._tc in self._fresh_cells:
            self._fresh_cells.discard(container._tc)
            return container.paragraphs[0]
        return container.add_paragraph()

    def _add_table(self, container, rows: int, cols: int):
        if isinstance(container, _Cell):
            self._fresh_cells.discard(container._tc)
            table = container.add_table(rows, cols)
        else:
            table = container.add_table(rows, cols)
        try:
            table.style = self.document.styles["Table Grid"]
        except KeyError:
            log.warning("Style 'Table Grid' missing from template")
        for row in table.rows:
            for cell in row.cells:
                self._fresh_cells.add(cell._tc)
        return table

    def _table(self, container, t: Table):
        if not t.rows or not t.rows[0]:
            return
        table = self._add_table(container, len(t.rows), len(t.rows[0]))
        for r, row in enumerate(t.rows):
            for c, cell in enumerate(row):
                if cell.row_span > 1 or cell.col_span > 1:
                    table.cell(r, c).merge(
                        table.cell(r + cell.row_span - 1, c + cell.col_span - 1)
                    )
        for r, row in enumerate(t.rows):
            if r < t.header_rows:
                table.rows[r]._tr.get_or_add_trPr().append(
                    parse_xml(f"<w:tblHeader {W_NS}/>")
                )
            for c, cell in enumerate(row):
                if cell.merged:
                    continue
                target = table.cell(r, c)
                self._blocks(target, cell.blocks)
                if r < t.header_rows:
                    for para in target.paragraphs:
                        for run in para.runs:
                            run.bold = True

    def _text_box(self, container, box: TextBox):
        table = self._add_table(container, 1, 1)
        cell = table.cell(0, 0)
        shading = f'<w:shd {W_NS} w:val="clear" w:color="auto" w:fill="F2F2F2"/>'
        cell._tc.get_or_add_tcPr().append(parse_xml(shading))
        self._blocks(cell, box.blocks)

    def _paragraph(self, container, p: Paragraph):
        if p.list_kind == "number":
            if self._number_list is None or p.restart:
                self._number_list = None
                self._number_formats = {}
            self._number_formats.setdefault(p.level, (p.number_format, p.number_suffix))
            if self._number_list is None:
                self._number_list = self.numbering.new_numbered_list(
                    self._number_formats
                )

        para = self._new_paragraph(container)
        if p.heading is not None:
            level = min(p.heading, 6)
            if level == 0:
                self._style(para, self.profile.style("title", "Title"))
            else:
                self._style(
                    para, self.profile.style(f"heading{level}", f"Heading {level}")
                )
        elif p.list_kind:
            self._style(para, self.profile.style("list", "List Paragraph"))
            num_id = (
                self.numbering.bullet_num
                if p.list_kind == "bullet"
                else self._number_list
            )
            _set_numbering(para, num_id, p.level)
        else:
            if self.profile.roles.get("body"):
                self._style(para, self.profile.roles["body"])
            if p.left_indent > 0:
                para.paragraph_format.left_indent = Pt(p.left_indent)

        if p.align in ALIGN:
            para.alignment = ALIGN[p.align]
        if p.page_break_before:
            para.paragraph_format.page_break_before = True
        self._items(para, p.items, heading=p.heading is not None)

    def _items(self, para, items, heading: bool = False, allow_images: bool = True):
        for item in items:
            if isinstance(item, Run):
                self._run(para, item, heading=heading)
            elif isinstance(item, Image) and allow_images:
                self._image(para, item)
            elif isinstance(item, Pencil) and allow_images:
                self._pencil(para, item)
            elif isinstance(item, PageBreak):
                para.add_run().add_break(WD_BREAK.PAGE)
            elif isinstance(item, Footnote):
                self._footnote(para, item)

    def _style(self, para, name: str):
        try:
            para.style = self.document.styles[name]
        except KeyError:
            log.warning("Style %r missing from template, using Normal", name)

    def _run(self, para, r: Run, heading: bool):
        run = para.add_run(r.text)
        if r.bold and not heading:
            run.bold = True
        if r.italic:
            run.italic = True
        if r.underline:
            run.underline = True
        if r.strike:
            run.font.strike = True
        if r.caps:
            run.font.all_caps = True
        if r.script == "super":
            run.font.superscript = True
        elif r.script == "sub":
            run.font.subscript = True
        if r.color and (self.heading_colors or not heading):
            run.font.color.rgb = RGBColor.from_string(r.color)
        if r.highlight:
            _set_shading(run, r.highlight)
        if r.lang and r.lang != self.language:
            _set_lang(run, r.lang)

    def _image(self, para, img: Image):
        width, height = img.width, img.height
        if width > self.max_image_width:
            height *= self.max_image_width / width
            width = self.max_image_width
        try:
            para.add_run().add_picture(
                io.BytesIO(img.data), width=Pt(width), height=Pt(height)
            )
        except Exception as exc:  # noqa: BLE001 - unsupported format (e.g. HEIC, PDF)
            log.warning("Could not embed image %s: %s", img.filename, exc)

    def _pencil(self, para, pencil: Pencil):
        img = pencil.image
        run = para.add_run()
        try:
            run.add_picture(
                io.BytesIO(img.data), width=Pt(img.width), height=Pt(img.height)
            )
        except Exception as exc:  # noqa: BLE001 - unsupported format
            log.warning("Could not embed pencil drawing %s: %s", img.filename, exc)
            return
        self._z += 1
        _anchor_picture(run, pencil.dx, pencil.dy, self._z)

    def _footnote(self, para, note: Footnote):
        note_id, element = self.footnotes.add()
        ref = para.add_run()
        ref._r.append(parse_xml(f'<w:footnoteReference {W_NS} w:id="{note_id}"/>'))
        self._char_style(ref, "Footnote Reference")
        paragraphs = [b for b in note.blocks if isinstance(b, Paragraph)] or [
            Paragraph()
        ]
        for i, p in enumerate(paragraphs):
            p_el = parse_xml(
                f'<w:p {W_NS}><w:pPr><w:pStyle w:val="FootnoteText"/></w:pPr></w:p>'
            )
            element.append(p_el)
            note_para = DocxParagraph(p_el, self.document._body)
            if i == 0:
                mark = note_para.add_run()
                mark._r.append(parse_xml(f"<w:footnoteRef {W_NS}/>"))
                self._char_style(mark, "Footnote Reference")
                note_para.add_run(" ")
            self._items(note_para, p.items, allow_images=False)

    def _char_style(self, run, name: str):
        try:
            run.style = self.document.styles[name]
        except KeyError:
            run.font.superscript = True


def write_docx(
    doc: Document,
    out: Path,
    template: Path | None = None,
    empty_lines_to_page_break: int = 4,
):
    WordWriter(doc, template, empty_lines_to_page_break).write(out)
