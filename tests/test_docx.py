from pathlib import Path

import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt
from test_rich import PNG

from pages2word.cli import main
from pages2word.docx_reader import DocxReader, _promote_numbered_headings
from pages2word.model import Cell, Document, Footnote, Image, Paragraph, Run, Table
from pages2word.styles import extract_styles
from pages2word.writer import write_docx

BODY = "Text odrážky, který je dost dlouhý, aby se jeho písmo počítalo do stylu. " * 2


def _bullets(document, glyph: str, left: int, font: str = "Symbol") -> str:
    """Add a one-level bullet list definition; returns its numId."""
    root = document.part.numbering_part.element
    num_id = str(len(root.findall(qn("w:num"))) + 1)
    abstract = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{num_id}">'
        f'<w:lvl w:ilvl="0"><w:numFmt w:val="bullet"/><w:lvlText w:val="{glyph}"/>'
        f'<w:pPr><w:ind w:left="{left}" w:hanging="360"/></w:pPr>'
        f'<w:rPr><w:rFonts w:ascii="{font}" w:hAnsi="{font}"/></w:rPr></w:lvl></w:abstractNum>'
    )
    nums = root.findall(qn("w:num"))
    (nums[0].addprevious if nums else root.append)(abstract)
    root.append(
        parse_xml(
            f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{num_id}"/></w:num>'
        )
    )
    return num_id


def _list_item(document, text: str, num_id: str, style: str = "List Paragraph"):
    p = document.add_paragraph(style=style)
    p._p.get_or_add_pPr().append(
        parse_xml(
            f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="{num_id}"/></w:numPr>'
        )
    )
    p.add_run(text)
    return p


def _heading_style(document, name: str, size: int):
    style = document.styles.add_style(name, docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = document.styles["Normal"]
    style.font.bold = True
    style.font.size = Pt(size)
    return style


def _my_styles(path: Path) -> Path:
    """The user's styled notes: custom headings, Calibri 12 set by hand on list text."""
    d = docx.Document()
    _heading_style(d, "Nadpis 1_M", 20)
    _heading_style(d, "Nadpis 2_M", 14)
    num = _bullets(d, "", 720)
    d.add_paragraph("1. Kapitola", style="Nadpis 1_M")
    d.add_paragraph("ODDÍL", style="Nadpis 2_M")
    for _ in range(3):
        run = _list_item(d, BODY, num).runs[0]
        run.font.name = "Calibri"
        run.font.size = Pt(12)
    d.save(path)
    return path


def _messy(path: Path) -> Path:
    """Notes pasted together from several sources."""
    d = docx.Document()
    _heading_style(d, "Nadpis 1_M", 16)
    helvetica = d.styles.add_style("Výchozí", docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)
    helvetica.font.name = "Helvetica Neue"
    arrows = _bullets(d, "", 284, "Wingdings")
    d.add_paragraph("2. Správní právo", style="Nadpis 1_M")
    item = _list_item(d, "Položka ve stylu Výchozí", arrows, style="Výchozí")
    item.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.JUSTIFY
    d.add_paragraph("• Ručně psaná odrážka")
    d.add_paragraph("• Další ručně psaná odrážka")
    d.add_paragraph("a) první možnost")
    d.add_paragraph("b) druhá možnost")
    d.add_paragraph("A. Novák uvádí, že…")
    d.save(path)
    return path


def _paragraphs(path: Path):
    return [p for p in docx.Document(str(path)).paragraphs if p.text.strip()]


def _num_id(p) -> str | None:
    ppr = p._p.pPr
    num = ppr.numPr if ppr is not None else None
    return str(num.numId.val) if num is not None and num.numId is not None else None


def test_extract_styles_stores_roles_and_folds_hand_formatting(tmp_path):
    out = extract_styles(_my_styles(tmp_path / "mine.docx"), tmp_path / "styles.docx")

    d = docx.Document(str(out))
    doc_vars = {
        v.get(qn("w:name")): v.get(qn("w:val"))
        for v in d.settings.element.iter(qn("w:docVar"))
    }
    assert doc_vars["pages2word.heading1"] == "Nadpis 1_M"
    assert doc_vars["pages2word.heading2"] == "Nadpis 2_M"
    assert doc_vars["pages2word.list"] == "List Paragraph"
    # Calibri 12 typed onto almost all text is now part of the base style
    assert d.styles["Normal"].font.name == "Calibri"
    assert d.styles["Normal"].font.size == Pt(12)
    # custom headings get an outline level, so the Navigation pane sees them
    outline = d.styles["Nadpis 1_M"].element.find(f"{qn('w:pPr')}/{qn('w:outlineLvl')}")
    assert outline is not None and outline.get(qn("w:val")) == "0"
    # the sample content uses the user's styles
    styles = {p.style.name for p in _paragraphs(out)}
    assert {"Nadpis 1_M", "Nadpis 2_M", "List Paragraph"} <= styles


def test_restyle_word_document_with_style_document(tmp_path):
    styles = extract_styles(
        _my_styles(tmp_path / "mine.docx"), tmp_path / "styles.docx"
    )
    messy = _messy(tmp_path / "messy.docx")

    assert main([str(messy), "-t", str(styles)]) == 0
    out = tmp_path / "messy (converted).docx"
    paras = _paragraphs(out)
    by_text = {p.text: p for p in paras}

    assert by_text["2. Správní právo"].style.name == "Nadpis 1_M"
    bullets = [
        by_text["Položka ve stylu Výchozí"],
        by_text["Ručně psaná odrážka"],
        by_text["Další ručně psaná odrážka"],
    ]
    assert {p.style.name for p in bullets} == {"List Paragraph"}
    assert len({_num_id(p) for p in bullets}) == 1  # one list, the user's bullets
    assert all(p.alignment is None for p in paras)  # hand-set justify is gone
    assert all(r.font.name is None for p in paras for r in p.runs)

    letters = [by_text["první možnost"], by_text["druhá možnost"]]
    assert _num_id(letters[0]) and _num_id(letters[0]) == _num_id(letters[1])
    assert by_text["A. Novák uvádí, že…"].style.name == "Normal"  # not a list


def test_template_without_stored_roles_is_learned(tmp_path):
    mine = _my_styles(tmp_path / "mine.docx")
    messy = _messy(tmp_path / "messy.docx")
    out = tmp_path / "out.docx"
    write_docx(DocxReader(messy).read(), out, mine)
    assert _paragraphs(out)[0].style.name == "Nadpis 1_M"


def test_input_document_is_never_overwritten(tmp_path):
    messy = _messy(tmp_path / "messy.docx")
    before = messy.read_bytes()
    assert main([str(messy), "--force"]) == 0
    assert messy.read_bytes() == before
    assert (tmp_path / "messy (converted).docx").exists()


def test_word_document_round_trip(tmp_path):
    src = tmp_path / "src.docx"
    doc = Document(
        [
            Paragraph([Run("Nadpis")], heading=1),
            Paragraph(
                [
                    Run("tučně", bold=True),
                    Run(" a modře", color="1F4E79"),
                    Footnote([Paragraph([Run("Poznámka")])]),
                ]
            ),
            Paragraph([Run("Číslo")], list_kind="number", restart=True),
            Paragraph([Run("Pod")], list_kind="bullet", level=1),
            Table(
                [
                    [Cell([Paragraph([Run("A")])], col_span=2), Cell(merged=True)],
                    [Cell([Paragraph([Run("B")])]), Cell([Paragraph([Run("C")])])],
                ],
                header_rows=1,
            ),
            Paragraph([Image(PNG, "x.png", 10, 10)]),
        ],
        "cs-CZ",
    )
    write_docx(doc, src)

    blocks = DocxReader(src).read().blocks
    heading, text, number, bullet, table, picture = blocks
    assert heading.heading == 1
    assert [(r.text, r.bold, r.color) for r in text.items if isinstance(r, Run)] == [
        ("tučně", True, None),
        (" a modře", False, "1F4E79"),
    ]
    note = next(i for i in text.items if isinstance(i, Footnote))
    assert note.blocks[0].text == "Poznámka"
    assert (number.list_kind, number.number_format, number.restart) == (
        "number",
        "decimal",
        True,
    )
    assert (bullet.list_kind, bullet.level) == ("bullet", 1)
    assert table.header_rows == 1
    assert table.rows[0][0].col_span == 2 and table.rows[0][1].merged
    assert table.rows[1][1].blocks[0].text == "C"
    assert isinstance(picture.items[0], Image) and picture.items[0].data == PNG


def test_numbered_titles_typed_as_text_join_their_headings():
    paras = [
        Paragraph([Run("22. Pravomoc")], heading=1),
        Paragraph([Run("23. Vláda a veřejná správa")]),
        Paragraph([Run("24. Správní úřady")]),
        Paragraph([Run("Běžný text")]),
        Paragraph([Run("7. Neodpovídá číslování")]),
        Paragraph([Run("1. Přímí vykonavatelé")]),
        Paragraph([Run("2. Nepřímí vykonavatelé")], heading=2),
    ]
    _promote_numbered_headings(paras)
    assert [p.heading for p in paras] == [1, 1, 1, None, None, 2, 2]
