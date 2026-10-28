from pathlib import Path

import docx

from pages2word.cli import main
from pages2word.model import Paragraph
from pages2word.pdf_reader import PdfReader, font_family, strip_layout

LONG = "This first line of the paragraph runs all the way to the right margin of the"


def _pdf(path: Path, lines: list[tuple[str, float, float, float, str]]) -> Path:
    """A one-page A4 PDF; lines are (font, size, x, y from top, text)."""
    fonts = {"F1": "Helvetica", "F2": "Helvetica-Bold", "F3": "Courier"}
    ops = []
    for font, size, x, top, text in lines:
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        ops.append(f"BT /{font} {size} Tf {x} {842 - top} Td ({escaped}) Tj ET")
    stream = "\n".join(ops).encode("cp1252")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R /Resources << /Font << "
        + b" ".join(f"/{k} {5 + i} 0 R".encode() for i, k in enumerate(fonts))
        + b" >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        *(
            f"<< /Type /Font /Subtype /Type1 /BaseFont /{name} /Encoding /WinAnsiEncoding >>".encode()
            for name in fonts.values()
        ),
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    path.write_bytes(bytes(out))
    return path


def _notes(path: Path) -> Path:
    return _pdf(
        path,
        [
            ("F2", 20, 200, 80, "Course notes"),
            ("F2", 14, 72, 120, "Introduction"),
            ("F1", 11, 72, 145, LONG),
            (
                "F1",
                11,
                72,
                159,
                "page, so the next line belongs to the same paragraph.",
            ),
            ("F1", 11, 72, 185, "A short paragraph."),
            ("F1", 11, 72, 205, "•"),
            ("F1", 11, 90, 205, "First bullet"),
            ("F3", 11, 108, 220, "o"),
            (
                "F1",
                11,
                126,
                220,
                "Nested bullet that is long enough to wrap onto a second",
            ),
            ("F1", 11, 126, 234, "line under its text"),
            ("F1", 11, 72, 250, "•"),
            ("F1", 11, 90, 250, "Second bullet"),
            ("F1", 10, 295, 800, "1"),
        ],
    )


def _paragraphs(doc) -> list[Paragraph]:
    return [b for b in doc.blocks if isinstance(b, Paragraph)]


def test_pdf_structure_is_rebuilt(tmp_path):
    doc = PdfReader(_notes(tmp_path / "notes.pdf")).read()
    paras = _paragraphs(doc)
    summary = [(p.heading, p.list_kind, p.level, p.text) for p in paras]
    assert summary == [
        (0, None, 0, "Course notes"),
        (1, None, 0, "Introduction"),
        (
            None,
            None,
            0,
            LONG + " page, so the next line belongs to the same paragraph.",
        ),
        (None, None, 0, "A short paragraph."),
        (None, "bullet", 0, "First bullet"),
        (
            None,
            "bullet",
            1,
            "Nested bullet that is long enough to wrap onto a second line under its text",
        ),
        (None, "bullet", 0, "Second bullet"),
    ]  # the page number is gone
    nested = paras[5]
    assert nested.left_indent == 126 - 72 and nested.first_line_indent == 108 - 126
    first_run = paras[2].items[0]
    assert (first_run.font, first_run.size, first_run.bold) == ("Helvetica", 11, False)
    assert doc.page is not None and doc.page.left == 72


def test_strip_layout_keeps_meaning_only(tmp_path):
    doc = PdfReader(_notes(tmp_path / "notes.pdf")).read()
    plain = strip_layout(doc)
    assert plain.page is None
    for p in _paragraphs(plain):
        assert (p.left_indent, p.first_line_indent, p.space_before, p.line_spacing) == (
            0,
            0,
            None,
            None,
        )
        assert all(r.font is None for r in p.items)
    assert _paragraphs(doc)[2].items[0].font == "Helvetica"  # original untouched


def test_pdf_gives_two_documents(tmp_path):
    pdf = _notes(tmp_path / "notes.pdf")
    assert main([str(pdf)]) == 0

    layout = docx.Document(str(tmp_path / "notes (PDF).docx"))
    body = next(p for p in layout.paragraphs if p.text.startswith("This first line"))
    assert body.runs[0].font.name == "Helvetica"
    assert body.runs[0].font.size.pt == 11
    assert round(layout.sections[0].left_margin.pt) == 72

    styled = docx.Document(str(tmp_path / "notes.docx"))
    assert all(r.font.name is None for p in styled.paragraphs for r in p.runs)
    bullets = [p for p in styled.paragraphs if p.style.name == "List Paragraph"]
    assert bullets[0].text == "First bullet"
    assert styled.paragraphs[1].style.name == "Heading 1"


def test_font_family_names():
    assert font_family("BCDEEE+TimesNewRomanPSMT") == "Times New Roman"
    assert font_family("ABCDEF+Calibri-BoldItalic") == "Calibri"
    assert font_family("CourierNewPSMT") == "Courier New"
    assert font_family("Arial,Bold") == "Arial"
