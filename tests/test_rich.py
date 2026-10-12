import struct
import zlib
from decimal import Decimal
from pathlib import Path

import docx
import pytest

from pages2word.cli import convert
from pages2word.model import (
    Cell,
    Document,
    Footnote,
    Image,
    Paragraph,
    Pencil,
    Run,
    Table,
    TextBox,
)
from pages2word.reader import PagesReader, _raw_fields, _reference, _split_out_blocks
from pages2word.tables import DECIMAL128_BIAS, decode_cell, format_number, split_row
from pages2word.writer import write_docx

FIXTURE = Path(__file__).parent / "fixtures" / "testPages2013.pages"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _png() -> bytes:
    def chunk(kind: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + kind
            + data
            + struct.pack(">I", zlib.crc32(kind + data))
        )

    header = struct.pack(">IIBBBBB", 1, 1, 8, 6, 0, 0, 0)  # 1x1 RGBA
    pixels = zlib.compress(bytes(5))  # filter byte + one transparent RGBA pixel
    signature = bytes.fromhex("89504e470d0a1a0a")
    return (
        signature
        + chunk(b"IHDR", header)
        + chunk(b"IDAT", pixels)
        + chunk(b"IEND", b"")
    )


PNG = _png()


def _decimal128(mantissa: int, exp: int) -> bytes:
    b = bytearray(mantissa.to_bytes(16, "little"))
    biased = exp + DECIMAL128_BIAS
    b[14] = (b[14] & 1) | ((biased & 0x7F) << 1)
    b[15] = (biased >> 7) & 0x7F
    return bytes(b)


def test_decode_v5_number_and_text():
    number = (
        bytes([5, 2, 0, 0, 0, 0, 0, 0]) + struct.pack("<I", 0x1) + _decimal128(125, -1)
    )
    assert decode_cell(number) == {"type": 2, "decimal": Decimal("12.5")}
    text = (
        bytes([5, 3, 0, 0, 0, 0, 0, 0])
        + struct.pack("<I", 0x8 | 0x20)
        + struct.pack("<ii", 7, 99)
    )
    assert decode_cell(text) == {"type": 3, "string": 7}


def test_split_row_with_empty_and_wide_offsets():
    buffer = b"A" * 8 + b"B" * 4
    offsets = struct.pack("<3h", 0, -1, 2)  # wide offsets are in 4-byte units
    assert split_row(buffer, offsets, 3, wide=True) == [b"A" * 8, None, b"B" * 4]


@pytest.mark.parametrize(
    ("value", "language", "expected"),
    [
        (Decimal("12.50"), "cs-CZ", "12,5"),
        (3.0, "en-US", "3"),
        (Decimal(1000), "cs-CZ", "1000"),
    ],
)
def test_format_number(value, language, expected):
    assert format_number(value, language) == expected


def test_reference_follows_nested_messages():
    # field 1 -> message{field 1 -> 4104}
    inner = bytes([0x08]) + bytes([0x88, 0x20])
    outer = bytes([0x0A, len(inner)]) + inner
    assert _reference(outer) == 4104
    assert _raw_fields(bytes([0x08, 0x05])) == [(1, 5)]


def test_table_inside_paragraph_becomes_block():
    table = Table([[Cell()]])
    p = Paragraph(items=[Run("before"), table, Run("after")], list_kind="bullet")
    blocks = _split_out_blocks(p)
    assert [type(b).__name__ for b in blocks] == ["Paragraph", "Table", "Paragraph"]
    assert blocks[0].list_kind == "bullet" and blocks[2].list_kind is None


def test_fixture_table_and_floating_text_box(tmp_path):
    doc = PagesReader(FIXTURE).read()
    tables = [b for b in doc.blocks if isinstance(b, Table)]
    assert len(tables) == 1
    assert [c.blocks[0].text for c in tables[0].rows[0]] == [
        "Column one",
        "Column two",
        "Column three",
    ]
    assert tables[0].header_rows == 1
    boxes = [b for b in doc.blocks if isinstance(b, TextBox)]
    assert boxes and boxes[0].blocks[0].text == "A text box with text."

    out = docx.Document(str(convert(FIXTURE, tmp_path / "out.docx")))
    assert len(out.tables) == 2  # table + text box, not merged into one
    assert out.tables[0].cell(3, 2).text == "Cell nine"


def test_writer_footnote_merge_and_pencil(tmp_path):
    table = Table(
        rows=[
            [Cell([Paragraph(items=[Run("Head")])], col_span=2), Cell(merged=True)],
            [Cell([Paragraph(items=[Run("A")])]), Cell([Paragraph(items=[Run("B")])])],
        ],
        header_rows=1,
    )
    pencil = Pencil(Image(PNG, "pencil.png", 20, 10), dx=5, dy=-2)
    doc = Document(
        [
            Paragraph(
                items=[
                    Run("Text"),
                    Footnote([Paragraph(items=[Run("Poznámka")])]),
                    pencil,
                ]
            ),
            table,
            TextBox([Paragraph(items=[Run("Box")])]),
        ],
        "cs-CZ",
    )
    path = tmp_path / "out.docx"
    write_docx(doc, path)
    out = docx.Document(str(path))

    body = out.element.body
    assert body.find(f".//{W}footnoteReference") is not None
    footnotes = [
        p
        for p in out.part.package.iter_parts()
        if str(p.partname) == "/word/footnotes.xml"
    ]
    assert footnotes and "Poznámka" in footnotes[0].blob.decode("utf-8")
    assert (
        body.find(".//{*}anchor") is not None
    )  # pencil floats instead of sitting inline
    assert out.tables[0].rows[0]._tr.find(f".//{W}gridSpan").get(f"{W}val") == "2"
    assert out.tables[0].rows[0]._tr.find(f".//{W}tblHeader") is not None
    assert len(out.tables) == 2
