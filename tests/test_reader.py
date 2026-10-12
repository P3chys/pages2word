from pathlib import Path

import docx
import pytest

from pages2word.cli import convert
from pages2word.reader import (
    Paragraph,
    Run,
    _convert_typed_bullets,
    _detect_headings,
    _normalize_list_levels,
    _Table,
)


def para(text: str, **kw) -> Paragraph:
    return Paragraph(
        items=[
            Run(text, **{k: v for k, v in kw.items() if k in Run.__dataclass_fields__})
        ],
        **{k: v for k, v in kw.items() if k in Paragraph.__dataclass_fields__},
    )


class _Entry:
    def __init__(self, index, value):
        self.character_index = index
        self.value = value


def test_para_style_table_carries_null_forward():
    entries = [_Entry(0, 5), _Entry(10, None), _Entry(20, 7), _Entry(30, None)]
    table = _Table(entries, lambda e: e.value, carry_forward=True)
    assert [table.at(i) for i in (0, 15, 25, 35)] == [5, 5, 7, 7]


@pytest.mark.parametrize(
    ("texts", "expected"),
    [
        (
            ["• A", "o B", "§ C"],
            [("bullet", 0, "A"), ("bullet", 1, "B"), ("bullet", 2, "C")],
        ),
        # "o" / "§ 4" as ordinary text outside a bullet run stay untouched
        (["o tom se nemluví"], [(None, 0, "o tom se nemluví")]),
        (
            ["• A", "§ 4 – Justiční spolupráce"],
            [("bullet", 0, "A"), (None, 0, "§ 4 – Justiční spolupráce")],
        ),
        # page number glued on by a PDF paste
        (["• A", "13o B"], [("bullet", 0, "A"), ("bullet", 1, "B")]),
        # hard-wrapped continuation line is merged back
        (
            ["• Rozhodnutí, které", "se týká věci"],
            [("bullet", 0, "Rozhodnutí, které se týká věci")],
        ),
    ],
)
def test_typed_bullets(texts, expected):
    paragraphs = [para(t) for t in texts]
    _convert_typed_bullets(paragraphs)
    assert [(p.list_kind, p.level, p.text) for p in paragraphs] == expected


def test_even_list_levels_are_compacted():
    paragraphs = [para("a", list_kind="bullet", level=lv) for lv in (0, 2, 4, 6, 2)]
    _normalize_list_levels(paragraphs)
    assert [p.level for p in paragraphs] == [0, 1, 2, 3, 1]


def test_heading_detection():
    paragraphs = [
        para("Body text " * 5, size=12),
        para("Big title", size=25, bold=True),
        para("Section", size=15, bold=True),
        para("Subsection", size=12, bold=True),
        para("Bold sentence that ends with a full stop.", size=12, bold=True),
    ]
    _detect_headings(paragraphs)
    assert [p.heading for p in paragraphs] == [None, 1, 2, 3, None]


SAMPLES = sorted(Path(__file__).parent.parent.glob("*.pages"))


@pytest.mark.skipif(not SAMPLES, reason="no sample .pages files")
@pytest.mark.parametrize("src", SAMPLES, ids=lambda p: p.name)
def test_convert_sample(src, tmp_path):
    out = convert(src, tmp_path / "out.docx")
    document = docx.Document(str(out))
    styles = {p.style.name for p in document.paragraphs}
    assert "List Paragraph" in styles
    assert "Heading 1" in styles
    lang = document.styles.element.find(".//{*}rPrDefault//{*}lang")
    assert (
        lang.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val")
        == "cs-CZ"
    )


def test_style_template_has_all_styles(tmp_path):
    from pages2word.styles import make_style_template

    document = docx.Document(str(make_style_template(tmp_path / "styly.docx")))
    names = {p.style.name for p in document.paragraphs}
    assert {
        "Title",
        "Heading 1",
        "Heading 2",
        "Heading 3",
        "List Paragraph",
        "Normal",
    } <= names


def test_unique_output_never_overwrites(tmp_path):
    from pages2word.app import unique_output

    src = tmp_path / "a.pages"
    assert unique_output(src).name == "a.docx"
    (tmp_path / "a.docx").touch()
    assert unique_output(src).name == "a (2).docx"
